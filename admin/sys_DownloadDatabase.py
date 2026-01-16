import streamlit as st
import admin.sys_sqlite as ss
import pandas as pd
from io import BytesIO
from openpyxl.utils import get_column_letter
import admin.sys_functions as fn
import time

def autofit_columns(sheet):
    for column_cells in sheet.columns:
        max_length = 0
        column = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column].width = adjusted_width

def app(selected):
    st.header("Tải xuống Cơ sở dữ liệu", divider='rainbow')
    st.info("Chức năng này sẽ xuất toàn bộ dữ liệu từ tất cả các bảng (trừ bảng hệ thống chứa 'sys') và các view ra một file Excel duy nhất.")

    if 'db_download_ready' not in st.session_state:
        st.session_state['db_download_ready'] = False

    if st.button("🚀 Bắt đầu Tải xuống", type="primary"):
        st.session_state['db_download_ready'] = False
        with st.spinner("Đang truy vấn và tạo file Excel... Vui lòng đợi, quá trình này có thể mất vài phút."):
            try:
                # 1. Get list of tables and views
                sql_tables = "SELECT name FROM sys.tables WHERE name NOT LIKE '%sys%' ORDER BY name"
                sql_views = "SELECT name FROM sys.views ORDER BY name"
                
                tables = ss.query_database_sqlite(sql_string=sql_tables, data_type='list')
                views = ss.query_database_sqlite(sql_string=sql_views, data_type='list')
                
                all_objects = tables + views
                
                # 2. Prepare Excel file in memory
                output = BytesIO()
                with pd.ExcelWriter(output, engine='openpyxl') as writer:
                    total_objects = len(all_objects)
                    progress_bar = st.progress(0, text="Đang xử lý...")

                    for i, obj_name in enumerate(all_objects):
                        try:
                            sheet_name = fn.sanitize_sheet_name(obj_name)
                            column_names = ss.get_table_columns(table_name=obj_name)
                            if not column_names:
                                raise Exception(f"Không thể lấy danh sách cột cho đối tượng '{obj_name}'.")
                            df = ss.get_columns_data(table_name=obj_name, columns=column_names)
                            df.to_excel(writer, sheet_name=sheet_name, index=False)
                            autofit_columns(writer.sheets[sheet_name])
                        except Exception as e:
                            error_df = pd.DataFrame({'Lỗi': [f"Không thể xuất dữ liệu từ '{obj_name}': {e}"]})
                            error_df.to_excel(writer, sheet_name=fn.sanitize_sheet_name(f"Loi_{obj_name}"), index=False)

                        progress_text = f"Đang xử lý: {obj_name} ({i+1}/{total_objects})"
                        progress_bar.progress((i + 1) / total_objects, text=progress_text)
                
                progress_bar.empty()
                st.session_state['db_download_data'] = output.getvalue()
                st.session_state['db_download_ready'] = True

            except Exception as e:
                st.error(f"Đã xảy ra lỗi trong quá trình tạo file: {e}")
                st.session_state['db_download_ready'] = False

    if st.session_state.get('db_download_ready', False):
        st.success("✅ File Excel đã sẵn sàng để tải xuống!")
        try:
            db_name = ss.get_database_name()
        except Exception:
            db_name = "Database"
        
        file_name = f"{db_name}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        st.download_button(
            label="📥 Tải xuống File Excel",
            data=st.session_state['db_download_data'],
            file_name=file_name,
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )